"""Extracts ethnicity, mother tongue, and religion for the 5 core cities
+ Alexandria from Romania's 2021 census DEFINITIVE results - the only
place this data exists (TEMPO-Online, the API used for every other INS
pull in this project, has NOTHING on ethnicity/religion/language-at-home,
confirmed by a full matrix-catalogue search).

Source: recensamantromania.ro's published result tables, downloaded once
to `data/location_features/census2021/*.xlsx` (Tables 2.02 ethnicity,
2.03 mother tongue, 2.04 religion - all at "Municipii, orase si comune",
i.e. locality, granularity, the finest level these are published at).

THREE caveats that matter more here than for any other pull in this
project:

1. This is a ONE-TIME 2021 census snapshot, not a refreshable annual
   series like everything else in this pipeline - there is no "pull the
   latest year" for this.
2. The census's own "resident population" figures are MEANINGFULLY LOWER
   than the administrative population figures already used everywhere
   else in this pipeline (from INS TEMPO-Online POP107D) - e.g. Bucharest
   1,716,961 (census-resident, 2021) vs 2,143,578 (administrative,
   POP107D, 2024). This is a REAL, known discrepancy in Romanian official
   statistics (registered/domiciled population vs. who was actually
   counted as physically resident), not a data error - do not silently
   mix these two population denominators.
3. Ethnicity/mother-tongue/religion are OPTIONAL, self-declared census
   questions. "Informatie nedisponibila" (not declared) is 15-25%+ of the
   population in every city checked here - a large share. Percentages
   below are computed against the CENSUS-RESIDENT total, so they include
   that non-response as its own category rather than silently excluding
   it - a "% Orthodox" figure that quietly dropped non-responders first
   would overstate every category.

Small counts ('*' in the source, an actual disclosure-control marker,
not a typo) are preserved as the string "<3 (suppressed)" rather than
guessed at or dropped - Romania's census tables suppress cell values
below a threshold specifically to prevent re-identification in
small-population breakdowns, standard statistical disclosure control.

Reproduce with: `PYTHONPATH=src python3 scripts/extract_census2021_demographics.py`
(requires the 3 xlsx files already downloaded to data/location_features/census2021/)
"""

import json

import openpyxl

CENSUS_DIR = "data/location_features/census2021"
OUT_FILE = "data/location_features/census2021_demographics.json"

TARGET_ROWS = {
    "Bucharest": "MUNICIPIUL BUCUREȘTI", "Brasov": "MUNICIPIUL BRAȘOV",
    "Calarasi": "MUNICIPIUL CĂLĂRAȘI", "Craiova": "MUNICIPIUL CRAIOVA",
    "Piatra-Neamt": "MUNICIPIUL PIATRA-NEAMȚ", "Alexandria": "MUNICIPIUL ALEXANDRIA",
}

TABLES = {
    "ethnicity": ("Tabel-2.02.1-si-Tabel-2.02.2.xlsx", "Tab 2.2.2"),
    "mother_tongue": ("Tabel-2.03.1-si-Tabel-2.03.2.xlsx", "Tab 2.3.2"),
    "religion": ("Tabel-2.04.1-si-Tabel-2.04.2.xlsx", "Tab 2.4.2"),
}


def extract_table(path: str, sheet: str, subheader_row: int = 4) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet]
    headers = [c.value for c in ws[subheader_row] if c.value is not None]

    results = {}
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        label = row[0]
        if not label:
            continue
        label_clean = str(label).strip()
        for city, match in TARGET_ROWS.items():
            if label_clean == match:
                raw_values = row[2 : 2 + len(headers)]
                breakdown = {}
                for header, val in zip(headers, raw_values):
                    if isinstance(val, (int, float)):
                        breakdown[header] = val
                    elif val == "*":
                        breakdown[header] = "<3 (suppressed)"
                    elif val == "-":
                        breakdown[header] = 0
                results[city] = {"census_resident_population_2021": row[1], "counts": breakdown}
    return results


def add_percentages(city_data: dict) -> None:
    total = city_data["census_resident_population_2021"]
    city_data["pct_of_resident_population"] = {
        k: round(100 * v / total, 2) for k, v in city_data["counts"].items() if isinstance(v, (int, float))
    }


def main():
    combined = {city: {} for city in TARGET_ROWS}
    for dimension, (filename, sheet) in TABLES.items():
        table = extract_table(f"{CENSUS_DIR}/{filename}", sheet)
        for city, data in table.items():
            add_percentages(data)
            combined[city][dimension] = data

    with open(OUT_FILE, "w") as f:
        json.dump(combined, f, indent=2, ensure_ascii=False)
    print(f"Wrote {OUT_FILE}")

    for city, dims in combined.items():
        pop = dims["ethnicity"]["census_resident_population_2021"]
        undeclared_eth = dims["ethnicity"]["pct_of_resident_population"].get("Informație nedisponibila", 0)
        print(f"{city}: census-resident pop {pop:,}, {undeclared_eth}% did not declare ethnicity")


if __name__ == "__main__":
    main()
